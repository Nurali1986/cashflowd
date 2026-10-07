"""Export the database back to an .xlsx with the same layout as the input sheets of the original workbook.

The file can be imported again (round trip), so it doubles as a backup.
"""
import io
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import logic, models
from .models import EXPENSE, INCOME, TRANSFER

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
OPS_HEADER = ["Qaysi sanada", " Qaysi oy uchun oyi (P&L)", "Turi (СТАТЬЯ)", "To'lov turi (СЧЕТ)", "Доллар",
              "Курс доллар", "Summa", "Kim uchun tolov (o'quvchi F.I.Sh)", "КОММЕНТАРИЙ\nCheck linki"]


def _header(ws, values):
    ws.append(values)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL


def _article_text(t: models.Transaction) -> str | None:
    if t.kind == TRANSFER:
        return f"ПЕРЕВОД. {t.account.name if t.account else ''} → {t.to_account.name if t.to_account else ''}"
    return t.article.full_name if t.article else None


def _operations(wb, title, txs):
    ws = wb.create_sheet(title)
    _header(ws, OPS_HEADER)
    ws.append(["НЕ УДАЛЯТЬ СТРОКУ!!"])
    for t in sorted(txs, key=lambda x: (x.date is None, x.date or x.pnl_month or date.min, x.id)):
        ws.append([
            t.date, logic.month_label(t.pnl_month) or None, _article_text(t),
            t.account.name if t.account and t.kind != TRANSFER else None,
            t.usd, t.usd_rate, t.amount, t.project.name if t.project else None, t.comment,
        ])
        ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"
        ws.cell(ws.max_row, 7).number_format = "#,##0"
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["H"].width = 30


def export_workbook(db: Session) -> bytes:
    settings = logic.get_settings(db)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    txs = list(db.scalars(select(models.Transaction)).unique())
    _operations(wb, "Cash flow", [t for t in txs if not t.is_plan])

    figures = logic.project_figures(db)
    ws = wb.create_sheet("P&L")
    _header(ws, ["ДАТА ЗАКАЗА", "Nima uchun (Proyekt)", "sinf (ЗАКАЗЧИК)", "ОПИСАНИЕ ЗАКАЗА", "КАТЕГОРИЯ",
                 "To'lov qilishi kerak (СТОИМОСТЬ ПРОЕКТА)", "СТАТУС", "ДАТА ЗАКРЫТИЯ", "ОПЛАЧЕНО\nФАКТ", "ДОЛГ",
                 "РАСХОДЫ", "МАРЖА", "РЕНТАБЕЛЬНОСТЬ, %", "КОММЕНТАРИЙ"])
    ws.append([])
    for p in db.scalars(select(models.Project).order_by(models.Project.sort, models.Project.id)):
        f = figures[p.id]
        ws.append([p.order_date, p.name, p.customer, p.description, p.category, p.contract_amount, p.status,
                   p.close_date, f["paid"], f["debt"], f["expenses"], f["margin"], f["profitability"], p.comment])
    ws.column_dimensions["B"].width = 32

    _operations(wb, "P&L Reja", [t for t in txs if t.is_plan])

    ws = wb.create_sheet("Kurs")
    _header(ws, ["Sana", "Dollar"])
    for r in db.scalars(select(models.ExchangeRate).order_by(models.ExchangeRate.date)):
        ws.append([r.date, r.rate])

    articles = logic.ordered_articles(db)
    accounts = logic.ordered_accounts(db)
    directory = list(db.scalars(select(models.DirectoryItem).order_by(models.DirectoryItem.sort)))
    ws = wb.create_sheet("справочник")
    _header(ws, ["ДАТА НАЧАЛА", None, "КАТЕГОРИЯ ДОХОДЫ", "КАТЕГОРИЯ (доходы)", "ПОДКАТЕГОРИЯ (доходы)", None,
                 "КАТЕГОРИЯ РАСХОДЫ", "КАТЕГОРИЯ (расходы)", "ПОДКАТЕГОРИЯ (расходы)", None, "СЧЕТ", None,
                 "СУММА НА НАЧАЛО", None, "ЗАКАЗЧИК", "КАТЕГОРИЯ ЗАКАЗА"])
    columns = {
        0: [logic.start_date(settings), None, "Учитывать в ОПУ дату закрытия проекта", None, None,
            "ДА" if settings.get("pnl_use_close_date") == "1" else "НЕТ"],
        2: list(dict.fromkeys(a.category for a in articles if a.type == INCOME)),
        3: [a.category for a in articles if a.type == INCOME and a.subcategory],
        4: [a.subcategory for a in articles if a.type == INCOME and a.subcategory],
        6: list(dict.fromkeys(a.category for a in articles if a.type == EXPENSE)),
        7: [a.category for a in articles if a.type == EXPENSE and a.subcategory],
        8: [a.subcategory for a in articles if a.type == EXPENSE and a.subcategory],
        10: [a.name for a in accounts],
        12: [a.opening_balance for a in accounts],
        14: [d.name for d in directory if d.kind == "customer"],
        15: [d.name for d in directory if d.kind == "order_category"],
    }
    for col, values in columns.items():
        for i, v in enumerate(values):
            ws.cell(row=i + 2, column=col + 1, value=v)

    ws = wb.create_sheet("настройки")
    ws.cell(3, 9, "для выпадающего списка")
    ws.cell(3, 14, "ДЛЯ СВОДА")
    for i, a in enumerate(articles):
        for j, v in enumerate([a.full_name, a.type, a.category, a.subcategory]):
            ws.cell(i + 4, 9 + j, v)
    row = 4
    for tree_row in logic.build_tree(articles):
        if tree_row.level == 1:
            values = [1, tree_row.type, None, tree_row.category]
        elif tree_row.label == logic.NO_SUBCATEGORY:
            continue
        else:
            values = [None, tree_row.type, tree_row.category, tree_row.label]
        for j, v in enumerate(values):
            ws.cell(row, 14 + j, v)
        row += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
