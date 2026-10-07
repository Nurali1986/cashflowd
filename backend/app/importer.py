"""Import of the «Pifagor - Cash flow - P&L» workbook into the database.

Only the input sheets are read (справочник, настройки tree, Cash flow, P&L Reja, P&L, Kurs);
every report sheet is recomputed by the web app.
"""
from datetime import date, datetime
from typing import Optional

import openpyxl
from sqlalchemy import delete
from sqlalchemy.orm import Session

from . import models
from .logic import parse_month_label
from .models import EXPENSE, INCOME, KINDS, TRANSFER


def _clean(value) -> Optional[str]:
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def _date(value) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _number(value) -> Optional[float]:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _rows(ws, min_row: int, max_col: int):
    for row in ws.iter_rows(min_row=min_row, max_col=max_col, values_only=True):
        yield tuple(row) + (None,) * (max_col - len(row))


def _article_key(kind, category, subcategory) -> tuple:
    return (kind, _clean(category), _clean(subcategory))


class _Importer:
    def __init__(self, db: Session):
        self.db = db
        self.articles: dict[tuple, models.Article] = {}
        self.accounts: dict[str, models.Account] = {}
        self.projects: dict[str, models.Project] = {}
        self.category_order: dict[tuple, int] = {}
        self.stats = {"transactions": 0, "plan": 0, "projects": 0, "articles": 0, "accounts": 0, "rates": 0,
                      "warnings": 0, "auto_projects": 0}

    # -- lookups that create missing entries ------------------------------------------
    def article(self, kind: str, category: Optional[str], subcategory: Optional[str]) -> Optional[models.Article]:
        key = _article_key(kind, category, subcategory)
        if kind not in (INCOME, EXPENSE) or not key[1]:
            return None
        if key not in self.articles:
            a = models.Article(type=kind, category=key[1], subcategory=key[2], sort=len(self.articles))
            self.db.add(a)
            self.articles[key] = a
            self.stats["articles"] += 1
        return self.articles[key]

    def account(self, name) -> Optional[models.Account]:
        name = _clean(name)
        if not name:
            return None
        if name not in self.accounts:
            acc = models.Account(name=name, opening_balance=0, sort=len(self.accounts))
            self.db.add(acc)
            self.accounts[name] = acc
            self.stats["accounts"] += 1
        return self.accounts[name]

    def project(self, name, auto: bool = False) -> Optional[models.Project]:
        name = _clean(name)
        if not name:
            return None
        if name not in self.projects:
            p = models.Project(name=name, sort=len(self.projects))
            if auto:
                p.comment = "Cash flow'dan avtomatik qo'shildi (Excel P&L ro'yxatida yo'q edi)"
                self.stats["auto_projects"] += 1
            self.db.add(p)
            self.projects[name] = p
        return self.projects[name]

    def parse_article(self, text):
        """«РАСХОД. Xodimlar maoshi. Oshpaz» → (kind, Article|None, from_account, to_account)."""
        text = _clean(text)
        if not text:
            return None, None, None, None
        parts = [p.strip() for p in text.split(".", 2)]
        kind = parts[0].upper()
        if kind not in KINDS:
            return None, None, None, None
        if kind == TRANSFER:
            rest = parts[1] + ("." + parts[2] if len(parts) > 2 else "") if len(parts) > 1 else ""
            src, _, dst = rest.partition("→")
            return kind, None, self.account(src), self.account(dst)
        category = parts[1] if len(parts) > 1 else None
        sub = parts[2] if len(parts) > 2 else None
        return kind, self.article(kind, category, sub), None, None

    # -- sheets ------------------------------------------------------------------------
    def settings_and_directories(self, wb):
        ws = wb["справочник"]
        rows = list(_rows(ws, 1, 21))
        start = _date(rows[1][0]) if len(rows) > 1 else None
        close_flag = _clean(rows[6][0]) if len(rows) > 6 else None
        settings = {
            "start_date": (start or date(2024, 7, 1)).isoformat(),
            "pnl_use_close_date": "1" if (close_flag or "").upper() == "ДА" else "0",
        }
        for key, value in settings.items():
            self.db.merge(models.Setting(key=key, value=value))

        customers, order_categories = [], []
        for r in rows[1:]:
            acc = self.account(r[10])
            if acc is not None and _number(r[12]) is not None:
                acc.opening_balance = _number(r[12])
            for target, value in ((customers, r[14]), (order_categories, r[15])):
                v = _clean(value)
                if v and v not in target:
                    target.append(v)
        for kind, items in (("customer", customers), ("order_category", order_categories)):
            for i, name in enumerate(items):
                self.db.add(models.DirectoryItem(kind=kind, name=name, sort=i))
        return rows

    def article_tree(self, wb, sprav_rows):
        ws = wb["настройки"]
        # «ДЛЯ СВОДА» (N:Q) defines the order of categories and subcategories.
        for r in _rows(ws, 4, 17):
            kind, category, sub = _clean(r[14]), _clean(r[15]), _clean(r[16])
            if kind in (INCOME, EXPENSE) and category and sub:
                self.article(kind, category, sub)
            elif kind in (INCOME, EXPENSE) and not category and sub:
                self.category_order.setdefault((kind, sub), len(self.category_order))
        # «для выпадающего списка» (I:L) — every article offered in the dropdown.
        for r in _rows(ws, 4, 12):
            kind, category, sub = _clean(r[9]), _clean(r[10]), _clean(r[11])
            if kind in (INCOME, EXPENSE) and category:
                self.article(kind, category, sub)
        # справочник category/subcategory pairs (D:E income, H:I expense).
        for r in sprav_rows[1:]:
            if _clean(r[3]) and _clean(r[4]):
                self.article(INCOME, r[3], r[4])
            if _clean(r[7]) and _clean(r[8]):
                self.article(EXPENSE, r[7], r[8])

    def finish_article_order(self):
        """Keep categories in the order of the настройки tree, articles inside in first-seen order."""
        items = sorted(self.articles.values(),
                       key=lambda a: (self.category_order.get((a.type, a.category), len(self.category_order)), a.sort))
        for i, a in enumerate(items):
            a.sort = i

    def projects_sheet(self, wb):
        for r in _rows(wb["P&L"], 3, 14):
            name = _clean(r[1])
            if not name:
                continue
            p = self.project(name)
            p.order_date = _date(r[0])
            p.customer = _clean(r[2])
            p.description = _clean(r[3])
            p.category = _clean(r[4])
            p.contract_amount = _number(r[5])
            status = _clean(r[6])
            p.status = status.lower() if status else None
            p.close_date = _date(r[7])
            p.comment = _clean(r[13])
            self.stats["projects"] += 1

    def rates(self, wb):
        seen = set()
        for r in _rows(wb["Kurs"], 2, 2):
            d, rate = _date(r[0]), _number(r[1])
            if d and rate and d not in seen:
                seen.add(d)
                self.db.add(models.ExchangeRate(date=d, rate=rate))
                self.stats["rates"] += 1

    def operations(self, wb, sheet: str, is_plan: bool):
        for r in _rows(wb[sheet], 3, 9):
            d = _date(r[0])
            has_content = any(v not in (None, "") for v in (r[1], r[2], r[3], r[7], r[8])) \
                or _number(r[4]) or r[6] not in (None, "", 0)
            if d is None and not has_content:
                continue
            warnings = []
            kind, article, src, dst = self.parse_article(r[2])
            if r[2] is not None and kind is None:
                warnings.append(f"Статья noma'lum: «{r[2]}»")
            elif kind is None:
                warnings.append("Статья ko'rsatilmagan — hisobotlarga kirmaydi")
            account = src if kind == TRANSFER else self.account(r[3])
            if account is None and kind != TRANSFER and d is not None:
                warnings.append("Счет ko'rsatilmagan — Kassa balansiga kirmaydi")

            amount = _number(r[6])
            if amount is None:
                if r[6] not in (None, ""):
                    warnings.append(f"Summa matn ko'rinishida edi: «{r[6]}» — 0 deb olindi, tuzating")
                elif _number(r[4]):
                    warnings.append("Dollar summasi bor, lekin so'mdagi summa hisoblanmagan")
                amount = 0.0

            pnl_raw = r[1]
            pnl_month = parse_month_label(pnl_raw)
            if pnl_raw and pnl_month is None:
                warnings.append(f"P&L oyi noto'g'ri: «{pnl_raw}»")
            elif not pnl_raw and kind in (INCOME, EXPENSE):
                warnings.append("P&L oyi ko'rsatilmagan — P&L hisobotiga kirmaydi")

            tx = models.Transaction(
                is_plan=is_plan,
                date=d,
                pnl_month=pnl_month,
                kind=kind,
                article=article,
                account=account,
                to_account=dst,
                usd=_number(r[4]),
                usd_rate=_number(r[5]),
                amount=amount,
                project=self.project(r[7], auto=True),
                comment=_clean(r[8]),
                import_warning="; ".join(warnings) or None,
            )
            self.db.add(tx)
            self.stats["plan" if is_plan else "transactions"] += 1
            if warnings:
                self.stats["warnings"] += 1


def import_workbook(db: Session, source) -> dict:
    """Replace all data with the content of the workbook (path or file object)."""
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    required = {"Cash flow", "P&L", "справочник"}
    missing = required - set(wb.sheetnames)
    if missing:
        raise ValueError(f"Faylda varaqlar topilmadi: {', '.join(sorted(missing))}")

    for table in (models.Receipt, models.Transaction, models.Project, models.Article, models.Account,
                  models.ExchangeRate, models.DirectoryItem):
        db.execute(delete(table))
    db.flush()

    imp = _Importer(db)
    sprav_rows = imp.settings_and_directories(wb)
    if "настройки" in wb.sheetnames:
        imp.article_tree(wb, sprav_rows)
    imp.projects_sheet(wb)
    if "Kurs" in wb.sheetnames:
        imp.rates(wb)
    imp.operations(wb, "Cash flow", is_plan=False)
    if "P&L Reja" in wb.sheetnames:
        imp.operations(wb, "P&L Reja", is_plan=True)
    imp.finish_article_order()
    db.commit()
    wb.close()
    return imp.stats
