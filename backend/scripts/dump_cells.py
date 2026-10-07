"""Print array formulas (text) and cached values for key cells."""
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

PATH = "/app/data_source.xlsx"
TARGETS = {
    "Cash flow": ["J2", "K2", "L2", "M2", "N2", "O2", "P2", "Q2", "R2", "S2", "T2", "U2", "V2", "W2", "X2", "Y2"],
    "P&L Reja": ["J2", "M2", "N2", "O2", "P2", "Q2", "R2"],
    "P&L": ["I3", "J3", "K3", "L3", "M3", "O3", "P3"],
    "Kassa": ["A3", "B3", "D3", "E3", "A1", "B1", "C1", "D1", "E1", "F1", "B2", "E2", "F5"],
    "Cash flow kunlik": ["A9", "B9", "C9", "D9", "E9", "F9", "A11", "B11", "C11", "D11", "E3", "F3", "F4", "E1"],
    "Cash flow oylik": ["F1", "A9", "B9", "C9", "D9", "E9", "F9", "F6", "F7", "F8", "F10", "E5", "E6"],
    "P&L oylik": ["A1", "B1", "C1", "D1", "E1", "F1", "G1", "A2", "E2", "F2", "G2", "A3", "B3", "E3", "F3", "G3", "A4", "E4", "F4", "G4", "A5", "E5", "F5", "G5", "A6", "B6", "C6", "D6", "E6", "F6", "G6", "A7", "D7", "G7"],
    "ОТЧЕТ по проекту": ["B5", "B6", "B7", "B8", "B9", "B10", "A11", "B11", "A12", "B12", "A13", "B13", "A14", "B14", "A15", "B15", "A16", "B16", "A17", "B17", "A18", "B18"],
    "Perevodlar": ["A1"],
    "справочник": ["R3", "A5", "A6", "A7"],
    "настройки": ["D3", "AI3", "AJ3", "M4", "R4", "X8", "Y8", "AP4", "AQ4", "AR4", "N5", "N6", "O5"],
    "настройки(дашбордОПУ)": ["C1", "W2", "X2", "Y2", "E6", "AN1", "AN6", "AN7", "AN9", "W15", "X15", "Y15"],
}

wb = openpyxl.load_workbook(PATH, read_only=False, data_only=False, keep_links=False) if False else None


def load(data_only):
    return openpyxl.load_workbook(PATH, read_only=True, data_only=data_only)


wf = load(False)
wv = load(True)


def cell_lookup(ws, coords):
    want = set(coords)
    out = {}
    maxr = max(int("".join(ch for ch in c if ch.isdigit())) for c in coords)
    for row in ws.iter_rows(min_row=1, max_row=maxr):
        for c in row:
            co = getattr(c, "coordinate", None)
            if co in want:
                out[co] = c.value
    return out


for sheet, coords in TARGETS.items():
    print("=" * 80)
    print("SHEET", sheet)
    f = cell_lookup(wf[sheet], coords)
    v = cell_lookup(wv[sheet], coords)
    for co in coords:
        fv = f.get(co)
        if isinstance(fv, ArrayFormula):
            fv = "ARRAY{" + str(fv.ref) + "} " + str(fv.text)
        print(f"  {co}: F={fv!r}\n        V={v.get(co)!r}")
