"""Dump full rows (formula + value) for given sheet/row ranges."""
import sys
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

PATH = "/app/data_source.xlsx"
SPEC = [
    ("P&L oylik", 1, 12, 30),
    ("Cash flow oylik", 1, 14, 17),
    ("Cash flow kunlik", 1, 2, 36),
    ("Cash flow kunlik", 9, 14, 8),
    ("Dashboard P&L", 1, 126, 28),
    ("Dashboard Cash flow", 1, 103, 28),
    ("настройки(дашбордДДС)", 1, 16, 54),
    ("настройки(дашбордОПУ)", 10, 16, 54),
]

wf = openpyxl.load_workbook(PATH, read_only=True, data_only=False)
wv = openpyxl.load_workbook(PATH, read_only=True, data_only=True)


def fmt(v):
    if isinstance(v, ArrayFormula):
        return "ARRAY{" + str(v.ref) + "} " + str(v.text)
    return v


for sheet, r1, r2, maxc in SPEC:
    print("=" * 90)
    print(sheet, r1, r2)
    fr = list(wf[sheet].iter_rows(min_row=r1, max_row=r2, max_col=maxc))
    vr = list(wv[sheet].iter_rows(min_row=r1, max_row=r2, max_col=maxc, values_only=True))
    for row_f, row_v in zip(fr, vr):
        for c, v in zip(row_f, row_v):
            f = getattr(c, "value", None)
            if f is None and v is None:
                continue
            co = getattr(c, "coordinate", "?")
            fs = fmt(f)
            if isinstance(fs, str) and len(fs) > 700:
                fs = fs[:700] + "..."
            if fs == v:
                print(f"  {co}: {v!r}")
            else:
                print(f"  {co}: F={fs!r} V={v!r}")
